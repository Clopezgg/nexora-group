from __future__ import annotations

"""Exportaciones financieras reales XLSX/PDF.

ORDEN MAESTRA DE CIERRE FINAL §19: los archivos son documentos nativos
(openpyxl/reportlab), nunca HTML/CSV renombrado. Los helpers son deliberadamente
genéricos para que todos los estados financieros compartan la misma fuente de
datos que sus endpoints JSON; no recalculan contabilidad en paralelo.
"""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(size=10, italic=True, color="555555")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = "1F3B57"
_TOTAL_FONT = Font(bold=True)
_MONEY_FORMAT = "#,##0.00"


def _header_block(sheet, *, title: str, company_name: str, currency_code: str, as_of: date) -> int:
    sheet["A1"] = title
    sheet["A1"].font = _TITLE_FONT
    sheet["A2"] = company_name
    sheet["A2"].font = Font(bold=True, size=11)
    sheet["A3"] = f"Moneda {currency_code} · Al {as_of.isoformat()}"
    sheet["A3"].font = _SUBTITLE_FONT
    sheet["A4"] = f"Generado {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}"
    sheet["A4"].font = _SUBTITLE_FONT
    return 6


def _write_columns(sheet, row: int, headers: Sequence[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")
        cell.alignment = Alignment(horizontal="center")


def _autosize(sheet, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        max_len = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 48)


def tabular_xlsx(
    *,
    title: str,
    company_name: str,
    currency_code: str,
    as_of: date,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    money_columns: set[int] | None = None,
    total_rows: Iterable[Sequence[object]] | None = None,
) -> bytes:
    """Construye un XLSX profesional sin perder precisión Decimal.

    ``money_columns`` usa índices base-cero. ``total_rows`` se imprime en
    negrita después del detalle.
    """
    wb = Workbook()
    sheet = wb.active
    sheet.title = title[:31]
    header_row = _header_block(
        sheet,
        title=title,
        company_name=company_name,
        currency_code=currency_code,
        as_of=as_of,
    )
    _write_columns(sheet, header_row, headers)
    money_columns = money_columns or set()

    current = header_row + 1
    for values in rows:
        for idx, value in enumerate(values):
            cell = sheet.cell(row=current, column=idx + 1, value=value)
            if idx in money_columns and value is not None:
                cell.number_format = _MONEY_FORMAT
        current += 1

    for values in total_rows or []:
        for idx, value in enumerate(values):
            cell = sheet.cell(row=current, column=idx + 1, value=value)
            cell.font = _TOTAL_FONT
            if idx in money_columns and value is not None:
                cell.number_format = _MONEY_FORMAT
        current += 1

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, current - 1)}"
    _autosize(sheet, len(headers))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def tabular_pdf(
    *,
    title: str,
    company_name: str,
    currency_code: str,
    as_of: date,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    total_rows: Iterable[Sequence[object]] | None = None,
) -> bytes:
    """PDF tabular auditable; no IDs técnicos ni datos inventados."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=title,
        author="NEXORA GROUP",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(company_name, styles["Heading3"]),
        Paragraph(
            f"Moneda {currency_code} · Al {as_of.isoformat()} · Generado {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}",
            styles["BodyText"],
        ),
        Spacer(1, 8),
    ]
    data = [list(headers)] + [["" if v is None else str(v) for v in row] for row in rows]
    totals = [["" if v is None else str(v) for v in row] for row in (total_rows or [])]
    data.extend(totals)
    table = Table(data, repeatRows=1)
    table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3B57")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d5dce3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if totals:
        table_style.extend(
            [
                ("FONTNAME", (0, -len(totals)), (-1, -1), "Helvetica-Bold"),
                ("LINEABOVE", (0, -len(totals)), (-1, -len(totals)), 0.8, colors.HexColor("#1F3B57")),
            ]
        )
    table.setStyle(TableStyle(table_style))
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def trial_balance_xlsx(
    *,
    company_name: str,
    currency_code: str,
    as_of: date,
    rows: list,
    total_debit: Decimal,
    total_credit: Decimal,
) -> bytes:
    return tabular_xlsx(
        title="Balance de Comprobación",
        company_name=company_name,
        currency_code=currency_code,
        as_of=as_of,
        headers=["Código", "Cuenta", "Debe", "Haber"],
        rows=[(row.account_code, row.account_name, row.debit_balance, row.credit_balance) for row in rows],
        money_columns={2, 3},
        total_rows=[("", "TOTAL", total_debit, total_credit)],
    )


def trial_balance_pdf(
    *,
    company_name: str,
    currency_code: str,
    as_of: date,
    rows: list,
    total_debit: Decimal,
    total_credit: Decimal,
) -> bytes:
    return tabular_pdf(
        title="Balance de Comprobación",
        company_name=company_name,
        currency_code=currency_code,
        as_of=as_of,
        headers=["Código", "Cuenta", "Debe", "Haber"],
        rows=[(row.account_code, row.account_name, row.debit_balance, row.credit_balance) for row in rows],
        total_rows=[("", "TOTAL", total_debit, total_credit)],
    )


def statement_xlsx(
    *,
    title: str,
    company_name: str,
    currency_code: str,
    as_of: date,
    sections: Sequence[tuple[str, Sequence]],
    totals: Sequence[tuple[str, Decimal]],
) -> bytes:
    rows: list[tuple[object, ...]] = []
    for section_name, section_rows in sections:
        rows.append((section_name, "", ""))
        rows.extend((row.account_code, row.account_name, row.balance) for row in section_rows)
    total_rows = [("", label, amount) for label, amount in totals]
    return tabular_xlsx(
        title=title,
        company_name=company_name,
        currency_code=currency_code,
        as_of=as_of,
        headers=["Código / Sección", "Cuenta", "Saldo"],
        rows=rows,
        money_columns={2},
        total_rows=total_rows,
    )


def statement_pdf(
    *,
    title: str,
    company_name: str,
    currency_code: str,
    as_of: date,
    sections: Sequence[tuple[str, Sequence]],
    totals: Sequence[tuple[str, Decimal]],
) -> bytes:
    rows: list[tuple[object, ...]] = []
    for section_name, section_rows in sections:
        rows.append((section_name, "", ""))
        rows.extend((row.account_code, row.account_name, row.balance) for row in section_rows)
    return tabular_pdf(
        title=title,
        company_name=company_name,
        currency_code=currency_code,
        as_of=as_of,
        headers=["Código / Sección", "Cuenta", "Saldo"],
        rows=rows,
        total_rows=[("", label, amount) for label, amount in totals],
    )
