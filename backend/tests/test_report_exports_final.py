import io
from decimal import Decimal

from openpyxl import load_workbook

from tests.helpers import create_account, create_company, login_admin


def _seed_balanced_posting(client):
    company = create_company(client, name="Exportadora NEXORA")
    cash = create_account(client, company_id=company["id"], code="1000", name="Caja", account_type="ASSET")
    revenue = create_account(client, company_id=company["id"], code="4000", name="Ingresos", account_type="REVENUE")
    response = client.post(
        "/api/accounting/journal-entries",
        json={
            "companyId": company["id"],
            "scope": "GENERAL",
            "currencyCode": "HNL",
            "lines": [
                {"accountId": cash["id"], "debitAmount": "125.50"},
                {"accountId": revenue["id"], "creditAmount": "125.50"},
            ],
        },
    )
    assert response.status_code == 201, response.text
    return company


def _assert_xlsx(response, expected_title: str):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content.startswith(b"PK")
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    assert workbook.active["A1"].value == expected_title


def _assert_pdf(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF")


def test_trial_balance_native_xlsx_and_pdf_are_real_documents(client):
    login_admin(client)
    company = _seed_balanced_posting(client)

    xlsx = client.get(f"/api/reports/trial-balance?companyId={company['id']}&format=xlsx")
    _assert_xlsx(xlsx, "Balance de Comprobación")
    workbook = load_workbook(io.BytesIO(xlsx.content), data_only=True)
    values = [cell.value for row in workbook.active.iter_rows() for cell in row]
    assert Decimal("125.50") in [Decimal(str(v)) for v in values if isinstance(v, (int, float, Decimal))]

    pdf = client.get(f"/api/reports/trial-balance?companyId={company['id']}&format=pdf")
    _assert_pdf(pdf)


def test_financial_statement_exports_cover_xlsx_and_pdf(client):
    login_admin(client)
    company = _seed_balanced_posting(client)

    cases = [
        ("balance-sheet", "Balance General"),
        ("income-statement", "Estado de Resultados"),
        ("cash-flow", "Flujo de Efectivo"),
        ("general-ledger", "Libro Mayor"),
        ("supplier-performance", "Desempeño de Proveedores"),
    ]
    for endpoint, title in cases:
        xlsx = client.get(f"/api/reports/{endpoint}/export?companyId={company['id']}&format=xlsx")
        _assert_xlsx(xlsx, title)
        _assert_pdf(client.get(f"/api/reports/{endpoint}/export?companyId={company['id']}&format=pdf"))


def test_budget_vs_actual_native_exports_reuse_real_budget_summary(client):
    login_admin(client)
    company = create_company(client, name="Presupuesto Export")
    project = client.post(
        "/api/projects",
        json={"companyId": company["id"], "name": "Casa Export", "code": "EXP-01", "currencyCode": "HNL"},
    )
    assert project.status_code == 201, project.text
    project_id = project.json()["id"]
    baseline = client.post(
        f"/api/projects/{project_id}/budgets/baseline",
        json={"currencyCode": "HNL", "lines": [{"authorizedAmount": "5000.00"}]},
    )
    assert baseline.status_code == 201, baseline.text

    xlsx = client.get(f"/api/reports/budget-vs-actual/export?projectId={project_id}&format=xlsx")
    _assert_xlsx(xlsx, "Presupuesto vs. Real — Casa Export")
    _assert_pdf(client.get(f"/api/reports/budget-vs-actual/export?projectId={project_id}&format=pdf"))


def test_report_export_rejects_unknown_format(client):
    login_admin(client)
    company = create_company(client)
    response = client.get(f"/api/reports/trial-balance?companyId={company['id']}&format=exe")
    assert response.status_code == 422
